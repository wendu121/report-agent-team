# tests/test_doc_render.py · 输出渲染（docx/pptx/pdf）hermetic 测试
# 覆盖：解析子集、三渲染器真实产出、PDF 中文字体防黑块、缺库响亮失败。

from __future__ import annotations

import io
import sys

import pytest

sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent.parent))

from tools.doc_render import (  # noqa: E402
    RendererUnavailable, parse_markdown, render_docx, render_pdf, render_pptx,
)

META = {"title": "锂电池行业研报", "topic": "锂电池", "generated_at": "2026-09-17 23:00", "task_id": "t1"}

SAMPLE_MD = """# 锂电池行业研报

## 摘要

固态电池产业化提速。头部厂商扩产明确。

## 风险

- 上游锂价波动剧烈
- 技术路线尚未收敛

| 厂商 | 产能 |
| --- | --- |
| A | 10GWh |
| B | 20GWh |

## 引用 / 来源

- [rec-1](https://a.com/1) — 标题一（可信度: high）
- [rec-2](https://a.com/2) — 标题二（可信度: medium）
"""


# ---------------------------------------------------------------- 解析
def test_parse_headings_and_levels():
    blocks = parse_markdown("# 大标题\n\n## 章节\n\n### 小节\n")
    heads = [b for b in blocks if b.kind == "heading"]
    assert [(h.level, h.text) for h in heads] == [(1, "大标题"), (2, "章节"), (3, "小节")]


def test_parse_bullets_grouped_into_one_block():
    blocks = parse_markdown("- 甲\n- 乙\n- 丙\n")
    bul = [b for b in blocks if b.kind == "bullets"]
    assert len(bul) == 1
    assert bul[0].items == ["甲", "乙", "丙"]


def test_parse_paragraph_and_unknown_line_degrades():
    blocks = parse_markdown("普通段落文本\n另一行并入同一段\n")
    paras = [b for b in blocks if b.kind == "paragraph"]
    assert len(paras) == 1
    assert "另一行并入同一段" in paras[0].text


def test_parse_table_skips_separator_row():
    blocks = parse_markdown(SAMPLE_MD)
    tb = [b for b in blocks if b.kind == "table"]
    assert len(tb) == 1
    assert tb[0].rows[0] == ["厂商", "产能"]
    # 分隔行 | --- | --- | 不得进入数据行
    assert all(not set("".join(r)) <= set("-: ") for r in tb[0].rows)
    assert ["A", "10GWh"] in tb[0].rows


def test_parse_citation_section_keeps_links():
    blocks = parse_markdown(SAMPLE_MD)
    cite_head = [b for b in blocks if b.kind == "heading" and "引用" in b.text]
    assert cite_head, "引用章节标题必须被解析出来"
    bul = [b for b in blocks if b.kind == "bullets" and any("rec-1" in i for i in b.items)]
    assert bul, "引用条目须保留 [id](url) 原样"


# ---------------------------------------------------------------- docx
def test_render_docx_produces_openable_document():
    data = render_docx(parse_markdown(SAMPLE_MD), META)
    assert isinstance(data, bytes) and len(data) > 1000
    from docx import Document
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "锂电池行业研报" in text
    assert "固态电池产业化提速" in text
    assert any(t.rows for t in doc.tables), "管道表须落成真表格"


def test_render_docx_hyperlink_for_citations():
    data = render_docx(parse_markdown(SAMPLE_MD), META)
    from docx import Document
    doc = Document(io.BytesIO(data))
    xml = doc.element.xml
    assert "w:hyperlink" in xml, "引用须生成真超链接元素（非裸文本）"
    # URL 不进 document.xml，而在关系部件 word/_rels/document.xml.rels
    targets = [r.target_ref for r in doc.part.rels.values() if r.is_external]
    assert any("a.com/1" in t for t in targets), "超链接目标须写入关系部件"


# ---------------------------------------------------------------- pptx
def test_render_pptx_cover_plus_section_slides():
    data = render_pptx(parse_markdown(SAMPLE_MD), META)
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) >= 4, "封面 + 摘要 + 风险 + 引用（+表页）"
    assert prs.slides[0].shapes.title.text == "锂电池行业研报"


def test_render_pptx_splits_overflow_bullets_into_continuation():
    md = "## 长章节\n" + "\n".join(f"- 要点{i}" for i in range(15))
    data = render_pptx(parse_markdown(md), META)
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    titles = [s.shapes.title.text for s in prs.slides if s.shapes.title]
    assert any("（续）" in t for t in titles), "超过 6 条须续页"
    # 封面 1 + 15 条按 6 分批 = 3 页
    assert sum(1 for t in titles if t.startswith("长章节")) == 3


def test_render_pptx_citation_slide_cap_is_ten():
    md = "## 引用 / 来源\n" + "\n".join(f"- [rec-{i}](https://a.com/{i}) — T{i}" for i in range(23))
    data = render_pptx(parse_markdown(md), META)
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    cite_slides = [s for s in prs.slides if s.shapes.title and "引用" in s.shapes.title.text]
    assert len(cite_slides) == 3, "23 条引用按每页 10 条 → 3 页"


# ---------------------------------------------------------------- pdf
def test_render_pdf_is_real_pdf_and_registers_cjk_font():
    data = render_pdf(parse_markdown(SAMPLE_MD), META)
    assert data[:4] == b"%PDF", "必须是真 PDF"
    # 防中文黑块：必须挂上 CID 字体（默认 Helvetica 不含 CJK）
    assert b"STSong-Light" in data, "reportlab 中文必须注册 UnicodeCIDFont，否则全黑块"


def test_render_pdf_contains_report_text():
    data = render_pdf(parse_markdown(SAMPLE_MD), META)
    # 内容流被压缩，只断言体积合理 + 含字体与链接目标（链接以 URI 形式存在于 PDF）
    assert len(data) > 2000
    assert b"a.com" in data, "引用链接须进入 PDF"


# ------------------------------------------------- 缺库必须响亮失败（诚实边界）
@pytest.mark.parametrize("fn,mod", [
    (render_docx, "docx"),
    (render_pptx, "pptx"),
    (render_pdf, "reportlab"),
])
def test_renderer_unavailable_raises_loudly(fn, mod, monkeypatch):
    monkeypatch.setitem(sys.modules, mod, None)  # 置 None → import 抛 ImportError
    with pytest.raises(RendererUnavailable):
        fn(parse_markdown(SAMPLE_MD), META)
